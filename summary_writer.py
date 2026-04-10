"""
Excel サマリー生成モジュール
openpyxl で invoices/summary.xlsx を作成・更新する。
"""

from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["日付", "件名", "送信者", "ファイル名", "金額（手動入力）"]
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")


def load_or_create_workbook(path: Path) -> Workbook:
    """既存ファイルをロード、なければヘッダー付きで新規作成する。"""
    if path.exists():
        return openpyxl.load_workbook(path)

    wb = Workbook()
    ws = wb.active
    ws.title = "請求書一覧"

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    return wb


def append_invoice_rows(workbook: Workbook, records: list[dict]) -> int:
    """
    新規レコードをサマリーに追記する。D列（ファイル名）で重複チェック。

    records の各要素:
        {
            "received_date": datetime,
            "subject": str,
            "sender_name": str,
            "sender_email": str,
            "filename": str,   # invoices/ 以下の相対パス
        }

    Returns: 追記した行数
    """
    ws = workbook.active

    # 既存ファイル名を収集（重複スキップ用）
    existing_filenames = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[3]:
            existing_filenames.add(str(row[3]))

    added = 0
    for rec in records:
        filename = rec["filename"]
        if filename in existing_filenames:
            continue

        sender = rec["sender_name"]
        if rec.get("sender_email"):
            sender = f"{rec['sender_name']} <{rec['sender_email']}>"

        row_data = [
            rec["received_date"].strftime("%Y-%m-%d"),
            rec["subject"],
            sender,
            filename,
            "",  # 金額は手動入力
        ]
        ws.append(row_data)
        existing_filenames.add(filename)
        added += 1

    _auto_adjust_columns(ws)
    return added


def save_summary(workbook: Workbook, path: Path) -> None:
    """ワークブックを保存する。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _auto_adjust_columns(ws) -> None:
    """列幅を内容に合わせて調整する。"""
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                value = str(cell.value) if cell.value else ""
                # 日本語文字は2文字分として計算
                length = sum(2 if ord(c) > 127 else 1 for c in value)
                if length > max_length:
                    max_length = length
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 60)
